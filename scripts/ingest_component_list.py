#!/usr/bin/env python3
"""
Bulk-ingest the VAC Standard Component List into the parts-knowledge DB as CATALOG
skeleton records. Only the authoritative columns (RIPPL code / model / make) are trusted;
everything else is left blank for later datasheet enrichment. Idempotent: existing codes
(active parts already enriched, or previously-ingested catalog rows) are skipped.

Run from ValveOS/:  python3 scripts/ingest_component_list.py
"""
import openpyxl, json, os, re

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(ROOT, "docs", "Datasheets", "VAC Standard Component List .xlsx")
DB   = os.path.join(ROOT, "data", "parts-knowledge.json")

# per-sheet column map: (code_col, model_col, make_col, category)
SHEETS = {
    "Ball Valve socket weld":     (0, 1, 2, "valve-ball-socketweld"),
    "Valmet Actuator":            (0, 1, 2, "actuator"),
    "Valmet LSB":                 (0, 2, 1, "limit-switch-box"),
    "SOV":                        (0, 2, 1, "sov"),
    "AFR":                        (0, 2, 1, "afr"),
    "MOR":                        (0, 2, 1, "gearbox-mor"),
    "POSITIONER":                 (0, 2, 1, "positioner"),
    "Ball valve flanged end":     (0, 1, 2, "valve-ball-flanged"),
    "MARSH ELECTRICAL ACTUATOR":  (0, 1, 2, "electric-actuator"),
    "BUTTERFLY VALVE PN10":       (0, 1, 2, "butterfly-valve-pn10"),
    "BUTTERFLY VALVE PN16":       (0, 1, 2, "butterfly-valve-pn16"),
}

CODE_RE = re.compile(r"^[A-Za-z]{2,4}\d{3,4}$")   # BV1121, PA1019, FR1001, MR1002, ...

def sv(c): return "" if c is None else str(c).strip()

def main():
    d = json.load(open(DB))
    parts = d["parts"]
    existing = {sv(p.get("itemSubCode")).upper() for p in parts.values() if p.get("itemSubCode")}
    existing |= {k.upper() for k in parts.keys()}

    wb = openpyxl.load_workbook(XLSX, data_only=True)
    added, per_cat = 0, {}
    for name, (cc, mc, kc, cat) in SHEETS.items():
        if name not in wb.sheetnames:
            continue
        for row in wb[name].iter_rows(values_only=True):
            if not row or cc >= len(row):
                continue
            code = sv(row[cc])
            if not CODE_RE.match(code):          # skips headers, SQ-only rows, blanks
                continue
            if code.upper() in existing:          # dedup vs active + already-ingested
                continue
            model = sv(row[mc]) if mc < len(row) else ""
            make  = sv(row[kc]) if kc < len(row) else ""
            parts[code] = {
                "partId": code,
                "itemSubCode": code,
                "displayName": (model or code) + (f" ({cat})" if cat else ""),
                "category": cat,
                "status": "catalog",
                "oem": {"make": make, "model": model,
                        "datasheetRef": "VAC Standard Component List (code/model/make authoritative)"},
                "interfaces": [], "tools": [], "fasteners": [], "rules": [], "checks": [], "hookups": [],
                "media": [],
                "provenance": {"lastUpdated": "2026-07-21",
                               "overallConfidence": "catalog skeleton - only RIPPL code / model / make are confirmed; specs to be enriched from a datasheet when this part is used in a config."},
            }
            existing.add(code.upper())
            added += 1
            per_cat[cat] = per_cat.get(cat, 0) + 1

    d.setdefault("meta", {})["catalogIngested"] = {
        "source": "docs/Datasheets/VAC Standard Component List .xlsx",
        "date": "2026-07-21",
        "note": "Every stocked RIPPL code seeded as a catalog skeleton (code/model/make only). status='catalog' until a config uses it, then it graduates to 'active' and gets datasheet-enriched.",
    }
    json.dump(d, open(DB, "w"), indent=2)
    total = len(parts)
    active = sum(1 for p in parts.values() if p.get("status") != "catalog")
    print(f"added {added} catalog skeletons across {len(per_cat)} categories")
    for c, n in sorted(per_cat.items(), key=lambda x: -x[1]):
        print(f"  {n:>3}  {c}")
    print(f"DB totals: {total} parts ({active} active, {total-active} catalog)")

if __name__ == "__main__":
    main()
